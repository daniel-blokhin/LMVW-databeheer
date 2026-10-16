"""Read plan records without losing physical source rows or inventing identifiers."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .settings import PipelineError, digest, file_sha256, text
from .settings import Config


@dataclass
class Delivery:
    delivery_id: str
    records: list[dict[str, Any]]
    geometry_by_key: dict[str, tuple[Any, str]]
    source_metadata: dict[str, Any]
    uses_geometry_file: bool


def _cell(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _read_rows(path: Path, *, sheet: str, delimiter: str) -> tuple[list[dict], list[str]]:
    rows = []
    if path.suffix.lower() == ".xlsx":
        values = load_workbook(path, read_only=True, data_only=True)
        formulas = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet not in values.sheetnames:
                raise PipelineError("Configured worksheet does not exist")
            ws, fs = values[sheet], formulas[sheet]
            if ws.max_row > 1_000_000:
                raise PipelineError("Workbook exceeds supported row count")
            header = [text(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            if any(v is None for v in header) or len(set(header)) != len(header):
                raise PipelineError("Workbook header has blank or duplicate column names")
            iterator = zip(ws.iter_rows(min_row=2, values_only=True), fs.iter_rows(min_row=2, values_only=True))
            for index, (row, formula_row) in enumerate(iterator, start=2):
                if all(v is None for v in formula_row):
                    continue
                record = {k: _cell(v) for k, v in zip(header, row)}
                missing_cache = [k for k, val, frm in zip(header, row, formula_row)
                                 if isinstance(frm, str) and frm.startswith("=") and val is None]
                formulae = {k: frm for k, frm in zip(header, formula_row)
                            if isinstance(frm, str) and frm.startswith("=")}
                rows.append({"source_row_number": index, "raw": record,
                             "missing_formula_cache": missing_cache, "formulae": formulae})
        finally:
            values.close()
            formulas.close()
    elif path.suffix.lower() == ".csv":
        csv.field_size_limit(100 * 1024 * 1024)
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, delimiter=delimiter)
            header = reader.fieldnames or []
            if not header or len(set(header)) != len(header) or any(not h for h in header):
                raise PipelineError("CSV header has blank or duplicate column names")
            for index, row in enumerate(reader, start=2):
                if None in row or any(value is None for value in row.values()):
                    raise PipelineError("Plan CSV row width differs from the header")
                if not any(value.strip() for value in row.values()):
                    continue
                rows.append({"source_row_number": index, "raw": row,
                             "missing_formula_cache": [], "formulae": {}})
    else:
        raise PipelineError("Plans must be an XLSX workbook or an explicitly delimited CSV")
    return rows, header


def read_delivery(config: Config) -> Delivery:
    settings = config.values["input"]
    path = config.resolve(settings["plans"])
    if not path.is_file():
        raise PipelineError("Plan input is missing; see input.plans in the configuration")
    plans_hash = file_sha256(path)
    sheet = settings.get("sheet", "Planlijst")
    records, headers = _read_rows(path, sheet=sheet, delimiter=settings.get("delimiter", ";"))
    if not records:
        raise PipelineError("Plan input contains no records")
    geometries: dict[str, tuple[Any, str]] = {}
    geo_hash = None
    geo_path_value = settings.get("geometries")
    if geo_path_value:
        import geopandas as gpd
        geo_path = config.resolve(geo_path_value)
        if not geo_path.is_file():
            raise PipelineError("Configured geometry file is missing; no fallback is used")
        geo_hash = file_sha256(geo_path)
        kwargs = {"layer": settings["geometry_layer"]} if settings.get("geometry_layer") else {}
        data = gpd.read_file(geo_path, **kwargs)
        if data.crs is None:
            raise PipelineError("Geometry file has no CRS; repair source metadata explicitly")
        key_field = settings.get("geometry_key", "Koppelcode")
        if key_field not in data:
            raise PipelineError("Geometry file has no configured join-key column")
        if file_sha256(geo_path) != geo_hash:
            raise PipelineError("Geometry input changed while being read")
        for _, row in data.iterrows():
            key = text(row[key_field])
            if not key or key in geometries:
                raise PipelineError("Geometry join keys are missing or non-unique; automatic dissolve is disabled")
            geometries[key] = (row.geometry, str(data.crs))
    if file_sha256(path) != plans_hash:
        raise PipelineError("Plan input changed while being read")
    delivery_id = digest({"plans_sha256": plans_hash, "geometry_sha256": geo_hash, "sheet": sheet})
    for record in records:
        record["record_id"] = digest([delivery_id, sheet, record["source_row_number"]])
    return Delivery(delivery_id, records, geometries, {
        "plans_sha256": plans_hash, "geometry_sha256": geo_hash,
        "sheet": sheet, "source_record_count": len(records), "source_columns": headers,
        "geometry_feature_count": len(geometries),
    }, bool(geo_path_value))
